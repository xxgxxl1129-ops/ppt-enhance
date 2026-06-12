"""主流水线: PDF → 解析 → 纠错 → 生成 → 评测."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from tqdm import tqdm

from ppt_enhance.agents.pipeline import CorrectionPipeline
from ppt_enhance.builder.pptx_builder import build_pptx
from ppt_enhance.config import settings
from ppt_enhance.eval.metrics import EvalReport, evaluate_conversion
from ppt_enhance.eval.renderer import pptx_to_images
from ppt_enhance.parser.docling_adapter import parse_with_docling
from ppt_enhance.parser.mineru_adapter import parse_with_mineru
from ppt_enhance.parser.pdf_renderer import render_pdf_pages
from ppt_enhance.parser.qwen_ocr_adapter import parse_with_qwen_ocr
from ppt_enhance.schemas.slide_ir import SlideIR


@dataclass
class PipelineResult:
    slide_ir: SlideIR
    pptx_path: Path
    eval_report: EvalReport | None = None
    correction_records: list = field(default_factory=list)
    work_dir: Path = Path(".")


def _build_outline_pptx(slide_ir: SlideIR, pptx_path: Path, progress_cb=None) -> int:
    """大纲逆推路线：每页用 VLM 逆推语义大纲，再用原生 PPT 元素重画。

    复用 outline_extractor + layout_engine（与 scripts/render_full.py 同逻辑，
    但改为实时逆推而非读缓存）。返回空白页数（无标题无节点）。
    extract_outline 自带降级：无 API key 时返回空 base outline，不抛错。

    VLM 逆推按页并行（OpenAI SDK 客户端线程安全，batch_outline.py 已验证），
    PPTX 写入串行并保持页序（python-pptx 非线程安全）。progress_cb(done, total)
    在每页逆推完成时回调，供 UI 显示实时进度。
    """
    from concurrent.futures import ThreadPoolExecutor, as_completed

    from pptx import Presentation

    from ppt_enhance.agents.llm_client import LLMClient
    from ppt_enhance.builder.layout_engine import render_outline_to_slide
    from ppt_enhance.parser.outline_extractor import extract_outline

    llm = LLMClient()
    pages = slide_ir.pages
    total = len(pages)

    def infer(idx, page):
        elems = [e for e in page.elements if e.is_textual and e.final_text.strip()]
        ol = extract_outline(page.render_path, page.page_no, page.width, page.height, elems, llm=llm)
        return idx, ol

    # 并行逆推（无 key 时 extract_outline 立即返回，并行无害）。
    # 进度回调只在主线程（as_completed 迭代处）触发，绝不在工作线程里
    # 碰任何 UI 框架——避免 Streamlit 的 NoSessionContext。
    workers = min(settings.vlm_concurrency, max(1, total))
    results: dict[int, object] = {}
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futures = [ex.submit(infer, i, p) for i, p in enumerate(pages)]
        for done_n, fut in enumerate(as_completed(futures), start=1):
            idx, ol = fut.result()
            results[idx] = ol
            if progress_cb:
                progress_cb(done_n, total)
    outlines = [results[i] for i in range(total)]  # 还原页序

    prs = Presentation()
    if prs.slides:  # 删默认空白页
        rid = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rid)
        del prs.slides._sldIdLst[0]

    empty_pages = 0
    for ol in outlines:  # 串行渲染，保持页序
        if not ol.nodes and not ol.title:
            empty_pages += 1
        render_outline_to_slide(prs, ol)

    prs.save(str(pptx_path))
    if not llm.available:
        print(
            "⚠️  outline 模式依赖 VLM(qwen3-vl) API key；当前未配置，"
            f"已生成 {empty_pages}/{total} 页空白版面。"
            "配置 .env 后重跑可得完整效果。"
        )
    return empty_pages


def run_pipeline(
    pdf_path: str | Path,
    output_dir: str | Path | None = None,
    mineru_json: str | Path | None = None,
    enable_correction: bool = True,
    enable_eval: bool = True,
    dpi: int | None = None,
    use_background: bool = True,
    ground_truth_text: str | None = None,
    parser: str = "docling",
    mode: str = "anchor",
    progress_cb=None,
) -> PipelineResult:
    pdf_path = Path(pdf_path)
    dpi = dpi or settings.default_dpi
    output_dir = Path(output_dir) if output_dir else pdf_path.parent / f"{pdf_path.stem}_output"
    work_dir = settings.work_dir / pdf_path.stem
    work_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    steps = ["解析 PDF", "智能纠错", "生成 PPTX", "质量评测"]

    def _report(frac: float, text: str) -> None:
        """progress_cb(0~1 进度, 描述)；UI 用它驱动进度条。无回调则忽略。"""
        if progress_cb:
            progress_cb(max(0.0, min(1.0, frac)), text)

    # 非 TTY 环境（Streamlit / 后台进程）下 stderr 管道可能已断，
    # tqdm 写入会触发 BrokenPipeError；此时禁用进度条（UI 有自己的进度展示）。
    import sys

    disable_bar = not sys.stderr.isatty()
    bar = tqdm(total=len(steps), desc="PPT Enhance", disable=disable_bar)

    # 1. 解析
    bar.set_description(steps[0])
    _report(0.02, f"解析 PDF（{parser}）…")
    if mineru_json:
        slide_ir = parse_with_mineru(pdf_path, mineru_json, work_dir, dpi=dpi)
    elif parser == "qwen-ocr":
        # OCR 阶段占进度 0.02→0.22，按页推进
        def _ocr_progress(d: int, t: int) -> None:
            _report(0.02 + 0.20 * (d / max(t, 1)), f"OCR 解析 {d}/{t} 页…")

        slide_ir = parse_with_qwen_ocr(pdf_path, work_dir, dpi=dpi, progress_cb=_ocr_progress)
    else:
        slide_ir = parse_with_docling(pdf_path, work_dir, dpi=dpi)
    ir_path = output_dir / "slide_ir.json"
    slide_ir.save(ir_path)
    bar.update(1)

    # 2. 纠错
    bar.set_description(steps[1])
    _report(0.25, "智能纠错…")
    pipeline = CorrectionPipeline()
    slide_ir = pipeline.run(slide_ir, enable_correction=enable_correction)
    slide_ir.save(output_dir / "slide_ir_corrected.json")
    bar.update(1)

    # 3. 生成
    bar.set_description(steps[2])
    pptx_path = output_dir / f"{pdf_path.stem}_enhanced.pptx"
    if mode == "outline":
        # 生成阶段占进度 0.35→0.85，按页推进
        def _outline_progress(d: int, t: int) -> None:
            _report(0.35 + 0.5 * (d / max(t, 1)), f"逆推大纲 {d}/{t} 页…")

        _build_outline_pptx(slide_ir, pptx_path, progress_cb=_outline_progress)
    else:
        _report(0.45, "生成 PPTX…")
        build_pptx(slide_ir, pptx_path, use_background=use_background)
    bar.update(1)

    # 4. 评测
    eval_report = None
    if enable_eval:
        bar.set_description(steps[3])
        _report(0.88, "质量评测（LibreOffice 渲染中）…")
        source_images = render_pdf_pages(pdf_path, work_dir / "eval_source", dpi=dpi)
        ppt_images, visual_reliable, rendered_pdf = pptx_to_images(pptx_path, work_dir / "eval_pptx", dpi=dpi)
        eval_report = evaluate_conversion(
            slide_ir,
            source_images,
            ppt_images,
            ground_truth_text=ground_truth_text,
            visual_reliable=visual_reliable,
            output_pdf=rendered_pdf,
        )
        import json
        (output_dir / "eval_report.json").write_text(
            json.dumps(eval_report.to_dict(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        bar.update(1)

    bar.set_description("完成")
    bar.close()
    _report(1.0, "完成")

    return PipelineResult(
        slide_ir=slide_ir,
        pptx_path=pptx_path,
        eval_report=eval_report,
        correction_records=pipeline.records,
        work_dir=work_dir,
    )
